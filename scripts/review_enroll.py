#!/usr/bin/env python3
"""Enroll a client's customer list into the Review Reactivation Engine.

Creates `contacts` (upsert-safe) + `review_requests` rows (status='pending',
step_number=1, campaign_type='reactivation') with next_send_at staggered at the
company's review pace. The dispatch-review-requests edge function does ALL
sending on its own cron — this script never sends anything.

Conventions copied from the app (AIMarketing.tsx) and the July Davis
enrollment:
  - contacts unique constraint: (client_id, phone); type='customer';
    new contacts get the 'Review Requested' tag (matches bulkAddTagToContacts).
  - tracking_slug: uuid4 hex, first 12 chars.
  - campaign_type 'reactivation' for list enrollments.
  - {first_name} greeting: when no human first name is derivable (condo
    associations, LLCs...), first_name is set to the literal 'there' so the
    SMS renders "Hey there, ..." — the dispatcher's own graceful fallback —
    instead of "Hey 136-146,".

Idempotent: rows whose phone (last-10 match) already has ANY review_requests
row for the company are skipped, so re-running never double-enrolls.

Sender preflight mirrors the dispatcher's resolution order (pin > own
approved subaccount > universal fallback + guard). If NO sender would
resolve, enrollment is refused.

Usage:
  set -a && source .env && set +a
  python3 scripts/review_enroll.py --slug quality-contracting-inc \
      --file /path/to/customerlist.xlsx --dry-run \
      [--name-cols contact,name] [--name-format auto|first-last|last-first-comma] \
      [--phone-cols "Main Phone"] [--email-col E-mail] \
      [--city-col City] [--state-col State/Province]

  --file also accepts "storage:<object-path>" to download from the Supabase
  'branding' storage bucket.

Column specifiers are header names (case-insensitive) or 0-based indexes.
"""

from __future__ import annotations

import argparse
import csv
import json
import os
import re
import sys
import uuid
from datetime import datetime, timedelta, timezone
from pathlib import Path
from zoneinfo import ZoneInfo

import requests

REPO_ROOT = Path(__file__).resolve().parent.parent
COMPANY_MAP = REPO_ROOT / "clients" / "company_map.json"

UNIVERSAL_FALLBACK_FROM = "+18337271056"  # HydroZ — sanctioned 2026-08-03

SEND_HOUR_START = 8   # dispatcher sends 08:00–17:59 company-local
SEND_HOUR_END_EXCL = 18

CHUNK = 500

BUSINESS_WORDS = {
    "llc", "inc", "inc.", "corp", "corp.", "co", "co.", "company", "trust",
    "condo", "condominium", "condominiums", "association", "assoc", "assoc.",
    "hoa", "properties", "property", "management", "mgmt", "realty", "dds",
    "pllc", "pc", "p.c.", "church", "school", "apartments", "homes",
    "builders", "group", "services", "service", "enterprises", "partners",
    "ltd", "ltd.", "lp", "llp", "bank", "hotel", "motel", "restaurant",
    "university", "college", "hospital", "clinic", "center", "centre",
    "dept", "department", "county", "city", "town", "village", "estates",
    "housing", "authority", "insurance", "agency", "solutions", "systems",
    "industries", "holdings", "ventures", "capital", "funeral", "storage",
}

HONORIFICS = {"mr", "mr.", "mrs", "mrs.", "ms", "ms.", "dr", "dr.", "miss",
              "prof", "prof.", "rev", "rev.", "fr", "fr.", "sir", "attn",
              "attn:"}

EMAIL_RE = re.compile(r"^[^@\s]+@[^@\s]+\.[a-zA-Z]{2,}$")
EXT_RE = re.compile(r"(?:ext\.?|x)\s*\d{1,6}\s*$", re.IGNORECASE)


# ---------------------------------------------------------------- helpers
def die(msg: str) -> None:
    print(f"ERROR: {msg}", file=sys.stderr)
    sys.exit(1)


class SB:
    """Tiny Supabase PostgREST client (service role)."""

    def __init__(self) -> None:
        self.url = os.environ.get("SUPABASE_URL", "").rstrip("/")
        self.key = os.environ.get("SUPABASE_SERVICE_ROLE_KEY", "")
        if not self.url or not self.key:
            die("SUPABASE_URL / SUPABASE_SERVICE_ROLE_KEY not set (source .env)")
        self.s = requests.Session()
        self.s.headers.update({
            "apikey": self.key,
            "Authorization": f"Bearer {self.key}",
        })

    def select(self, table: str, params: dict) -> list:
        out, offset, page = [], 0, 1000
        while True:
            p = dict(params)
            p["limit"] = page
            p["offset"] = offset
            r = self.s.get(f"{self.url}/rest/v1/{table}", params=p, timeout=60)
            r.raise_for_status()
            batch = r.json()
            out.extend(batch)
            if len(batch) < page:
                return out
            offset += page

    def insert(self, table: str, rows: list, upsert_on: str | None = None,
               returning: bool = False) -> list:
        prefer = ["return=representation" if returning else "return=minimal"]
        params = {}
        if upsert_on:
            prefer.append("resolution=merge-duplicates")
            params["on_conflict"] = upsert_on
        got = []
        for i in range(0, len(rows), CHUNK):
            chunk = rows[i:i + CHUNK]
            r = self.s.post(
                f"{self.url}/rest/v1/{table}", params=params,
                headers={"Prefer": ",".join(prefer),
                         "Content-Type": "application/json"},
                data=json.dumps(chunk), timeout=120)
            if r.status_code >= 300:
                die(f"insert into {table} failed ({r.status_code}): {r.text[:500]}")
            if returning:
                got.extend(r.json())
        return got

    def download_branding(self, object_path: str, dest: Path) -> None:
        r = self.s.get(f"{self.url}/storage/v1/object/branding/{object_path}",
                       timeout=120)
        if r.status_code != 200:
            die(f"storage download failed ({r.status_code}) for {object_path}")
        dest.write_bytes(r.content)


def normalize_phone(raw) -> str | None:
    """To +1E164; None when not a valid NANP number."""
    if raw is None:
        return None
    s = str(raw).strip()
    if not s:
        return None
    s = EXT_RE.sub("", s)
    digits = re.sub(r"\D", "", s)
    if len(digits) == 11 and digits.startswith("1"):
        digits = digits[1:]
    if len(digits) != 10:
        return None
    if digits[0] not in "23456789" or digits[3] not in "23456789":
        return None
    return "+1" + digits


def clean_email(raw) -> str | None:
    if raw is None:
        return None
    s = str(raw).strip().strip(",;").strip()
    return s if s and EMAIL_RE.match(s) else None


def fix_caps(token: str) -> str:
    return token.title() if token.isupper() and len(token) > 2 else token


def looks_like_person(full_name: str) -> bool:
    low = full_name.lower()
    if any(ch.isdigit() for ch in full_name):
        return False
    tokens = re.split(r"[\s,/&]+", low)
    return not any(t.strip(".,()") in BUSINESS_WORDS for t in tokens if t)


def derive_first_name(person_part: str, full_name: str) -> str:
    """First name for the SMS greeting, or 'there' when not derivable."""
    if not looks_like_person(full_name):
        return "there"
    tokens = [t for t in re.split(r"\s+", person_part.strip()) if t]
    tokens = [t for t in tokens if t.lower().strip(".") not in HONORIFICS]
    if not tokens:
        return "there"
    first = fix_caps(tokens[0]).strip(".,&/")
    if len(first) < 2 or not re.match(r"^[A-Za-z][A-Za-z'.-]*$", first):
        return "there"
    return first


# ---------------------------------------------------------------- file IO
def load_rows(path: Path) -> tuple[list, list]:
    """Returns (header, data_rows) as lists of str."""
    if path.suffix.lower() in (".xlsx", ".xlsm"):
        import openpyxl
        wb = openpyxl.load_workbook(path, read_only=True)
        ws = wb[wb.sheetnames[0]]
        rows = [["" if c is None else str(c).strip() for c in r]
                for r in ws.iter_rows(values_only=True)]
        wb.close()
    else:
        with open(path, newline="", encoding="utf-8-sig", errors="replace") as f:
            rows = [[c.strip() for c in r] for r in csv.reader(f)]
    rows = [r for r in rows if any(c for c in r)]
    if not rows:
        die(f"{path} is empty")
    return rows[0], rows[1:]


def resolve_col(spec: str, header: list) -> int:
    spec = spec.strip()
    if re.fullmatch(r"\d+", spec):
        idx = int(spec)
        if idx >= len(header):
            die(f"column index {idx} out of range (header has {len(header)})")
        return idx
    low = [h.lower() for h in header]
    if spec.lower() in low:
        return low.index(spec.lower())
    die(f"column '{spec}' not found in header {header}")


def auto_detect(header: list, data: list, kind: str) -> list:
    """Header-name candidates validated against the data; falls back to a
    pure data scan (handles files whose data is shifted vs the header)."""
    n = max(len(r) for r in data)

    def ratio(idx, pred):
        vals = [r[idx] for r in data if idx < len(r) and r[idx]]
        if not vals:
            return 0.0
        return sum(1 for v in vals if pred(v)) / len(vals)

    if kind == "phone":
        pred = lambda v: normalize_phone(v) is not None
        head_hits = [i for i, h in enumerate(header)
                     if re.search(r"mobile|cell", h, re.I)] + \
                    [i for i, h in enumerate(header)
                     if re.search(r"phone", h, re.I) and not re.search(r"fax", h, re.I)]
    else:  # email
        pred = lambda v: clean_email(v) is not None
        head_hits = [i for i, h in enumerate(header) if re.search(r"e-?mail", h, re.I)]

    good = [i for i in dict.fromkeys(head_hits) if ratio(i, pred) >= 0.5]
    if good:
        return good
    return [i for i in range(n) if ratio(i, pred) >= 0.5]


# ---------------------------------------------------------------- schedule
def stagger_times(count: int, tz_name: str, pace_per_20min: int,
                  start: datetime | None = None) -> list:
    tz = ZoneInfo(tz_name or "America/New_York")
    interval = timedelta(minutes=20.0 / max(pace_per_20min, 1))
    cursor = (start or datetime.now(timezone.utc)).astimezone(tz)
    out = []
    for _ in range(count):
        if cursor.hour < SEND_HOUR_START:
            cursor = cursor.replace(hour=SEND_HOUR_START, minute=0, second=0,
                                    microsecond=0)
        elif cursor.hour >= SEND_HOUR_END_EXCL:
            cursor = (cursor + timedelta(days=1)).replace(
                hour=SEND_HOUR_START, minute=0, second=0, microsecond=0)
        out.append(cursor.astimezone(timezone.utc)
                   .isoformat(timespec="milliseconds").replace("+00:00", "Z"))
        cursor += interval
    return out


# ---------------------------------------------------------------- sender
def sender_preflight(sb: SB, company_id: str, fallback_from: str) -> tuple[str | None, str]:
    """Mirror the dispatcher's resolution. Returns (sender_phone, mode) or
    (None, reason)."""
    setups = sb.select("company_phone_setup", {
        "id": f"eq.{company_id}",
        "select": "id,compliance_status,twilio_subaccount_sid,twilio_auth_token,agent_phone_1,review_sender",
    })
    setup = setups[0] if setups else None

    pin = (setup or {}).get("review_sender") or None
    if isinstance(pin, dict) and pin.get("account_sid") and pin.get("from"):
        if pin.get("auth_mode") == "subaccount" and setup and \
                setup.get("twilio_subaccount_sid") == pin["account_sid"] and \
                setup.get("twilio_auth_token"):
            return pin["from"], "pinned subaccount"
        if pin.get("auth_mode") == "master_fallback":
            return pin["from"], "pinned master_fallback"

    if setup and setup.get("compliance_status") == "approved" and \
            setup.get("twilio_subaccount_sid") and setup.get("twilio_auth_token") and \
            setup.get("agent_phone_1"):
        return setup["agent_phone_1"], "own approved subaccount"

    # Universal fallback + the dispatcher's ownership guard: refused if the
    # fallback number matches ANY company_phone_setup row (a current client).
    owned = sb.select("company_phone_setup", {
        "agent_phone_1": f"eq.{fallback_from}", "select": "id"})
    if owned:
        return None, (f"fallback {fallback_from} is owned by client "
                      f"{owned[0]['id']} — dispatcher guard would refuse it")
    return fallback_from, "universal fallback (master creds)"


# ---------------------------------------------------------------- main
def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("--slug", required=True)
    ap.add_argument("--file", required=True,
                    help="local path, or storage:<object-path> in the branding bucket")
    ap.add_argument("--dry-run", action="store_true")
    ap.add_argument("--name-cols", default=None,
                    help="ordered, comma-separated; first non-empty wins")
    ap.add_argument("--name-format", default="auto",
                    choices=["auto", "first-last", "last-first-comma"])
    ap.add_argument("--phone-cols", default=None,
                    help="ordered, comma-separated; first valid number wins")
    ap.add_argument("--email-col", default=None)
    ap.add_argument("--city-col", default=None)
    ap.add_argument("--state-col", default=None)
    ap.add_argument("--fallback-from", default=UNIVERSAL_FALLBACK_FROM)
    ap.add_argument("--staged", action="store_true",
                    help="write rows as status='staged' — the dispatcher "
                         "never touches them until the app's Activate "
                         "Review Campaign button (or a manual UPDATE) flips "
                         "them to 'pending' (Santino 2026-08-23)")
    args = ap.parse_args()

    sb = SB()

    company_map = json.loads(COMPANY_MAP.read_text())
    company_id = company_map.get(args.slug)
    if not company_id:
        cfg = REPO_ROOT / "clients" / f"{args.slug}.json"
        if cfg.exists():
            company_id = json.loads(cfg.read_text()).get("company_id")
    if not company_id:
        die(f"no company_id for slug '{args.slug}'")

    companies = sb.select("companies", {
        "id": f"eq.{company_id}",
        "select": "id,name,status,timezone,account_owner_name,"
                  "review_pace_per_20min,google_review_url,review_gate_enabled"})
    if not companies:
        die(f"company {company_id} not found")
    co = companies[0]
    if str(co.get("status", "")).strip().lower() in (
            "paused", "suspended", "cancelled", "canceled", "churned", "inactive", "archived"):
        die(f"company status is '{co['status']}' — dispatcher would never send")

    sender, sender_mode = sender_preflight(sb, company_id, args.fallback_from)
    if not sender:
        die(f"NO SENDER would resolve for {co['name']}: {sender_mode} — not enrolling")

    # ---- file ----
    if args.file.startswith("storage:"):
        obj = args.file[len("storage:"):]
        local = Path("/tmp") / Path(obj).name
        sb.download_branding(obj, local)
    else:
        local = Path(args.file)
        if not local.exists():
            die(f"file not found: {local}")
    header, data = load_rows(local)

    name_cols = ([resolve_col(c, header) for c in args.name_cols.split(",")]
                 if args.name_cols else [0])
    phone_cols = ([resolve_col(c, header) for c in args.phone_cols.split(",")]
                  if args.phone_cols else auto_detect(header, data, "phone"))
    if not phone_cols:
        die("no phone column found — pass --phone-cols")
    email_cols = ([resolve_col(args.email_col, header)]
                  if args.email_col else auto_detect(header, data, "email"))
    city_col = (resolve_col(args.city_col, header) if args.city_col else
                next((i for i, h in enumerate(header) if h.lower() == "city"), None))
    state_col = (resolve_col(args.state_col, header) if args.state_col else
                 next((i for i, h in enumerate(header)
                       if h.lower() in ("state", "state/province")), None))

    def cell(row, idx):
        return row[idx].strip() if idx is not None and idx < len(row) else ""

    # name format autodetect
    name_format = args.name_format
    if name_format == "auto":
        vals = [cell(r, name_cols[0]) for r in data if cell(r, name_cols[0])]
        commaish = sum(1 for v in vals
                       if re.match(r"^[^,\d]+,\s*[^,\d]+$", v))
        name_format = ("last-first-comma"
                       if vals and commaish / len(vals) >= 0.6 else "first-last")

    # ---- parse ----
    parsed, counts = [], {"rows_in_file": len(data), "no_phone": 0,
                          "dup_in_list": 0, "generic_greeting": 0}
    seen = set()
    for row in data:
        phone = None
        for pc in phone_cols:
            phone = normalize_phone(cell(row, pc))
            if phone:
                break
        if not phone:
            counts["no_phone"] += 1
            continue
        if phone in seen:
            counts["dup_in_list"] += 1
            continue
        seen.add(phone)

        raw_name = next((cell(row, nc) for nc in name_cols if cell(row, nc)), "")
        if name_format == "last-first-comma" and "," in raw_name:
            last_part, first_part = [p.strip() for p in raw_name.split(",", 1)]
            display = f"{first_part} {last_part}".strip()
        else:
            all_toks = raw_name.split()
            toks = [t for t in all_toks
                    if t.lower().strip(".:") not in HONORIFICS]
            # "Dr. Gross" (honorific + single token) is a LAST name — no
            # usable first name, keep the graceful greeting.
            if len(toks) < len(all_toks) and len(toks) == 1:
                toks = []
            first_part = toks[0] if toks else ""
            last_part = " ".join(toks[1:])
            display = raw_name
        display = " ".join(fix_caps(t) for t in display.split()) or "Customer"
        first_name = derive_first_name(first_part, raw_name or display)
        if first_name == "there":
            counts["generic_greeting"] += 1
            last_name = None
        else:
            last_name = " ".join(fix_caps(t) for t in last_part.split()) or None

        email = next((clean_email(cell(row, ec)) for ec in email_cols
                      if clean_email(cell(row, ec))), None)
        parsed.append({
            "name": display, "first_name": first_name, "last_name": last_name,
            "phone": phone, "email": email,
            "city": cell(row, city_col) or None,
            "state": cell(row, state_col) or None,
        })

    # ---- existing state ----
    existing_contacts = sb.select("contacts", {
        "client_id": f"eq.{company_id}",
        "select": "id,phone,opted_out"})
    existing_requests = sb.select("review_requests", {
        "company_id": f"eq.{company_id}",
        "select": "contact_id,status,opted_out"})

    by_last10 = {}
    for c in existing_contacts:
        d = re.sub(r"\D", "", c.get("phone") or "")[-10:]
        if len(d) == 10:
            by_last10.setdefault(d, []).append(c)
    contacts_with_request = {r["contact_id"] for r in existing_requests}
    optout_contact_ids = {c["id"] for c in existing_contacts if c.get("opted_out")} | \
                         {r["contact_id"] for r in existing_requests
                          if r.get("opted_out") or r.get("status") == "unsubscribed"}

    to_enroll = []
    counts["opted_out"] = counts["already_enrolled"] = 0
    for p in parsed:
        matches = by_last10.get(p["phone"][-10:], [])
        if any(m["id"] in optout_contact_ids for m in matches):
            counts["opted_out"] += 1
            continue
        if any(m["id"] in contacts_with_request for m in matches):
            counts["already_enrolled"] += 1
            continue
        p["existing_contact_id"] = matches[0]["id"] if matches else None
        to_enroll.append(p)
    counts["to_enroll"] = len(to_enroll)

    pace = co.get("review_pace_per_20min") or 1
    schedule = stagger_times(len(to_enroll), co.get("timezone"), pace)

    # ---- report header ----
    print(f"== {co['name']} ({company_id}) — {args.slug} ==")
    print(f"file: {local.name} | rows: {counts['rows_in_file']} | "
          f"name-format: {name_format}")
    print(f"sender that will resolve: {sender} ({sender_mode})")
    print(f"google_review_url: {co.get('google_review_url') or 'MISSING'}")
    print(f"review_gate_enabled: {co.get('review_gate_enabled')} | "
          f"pace: {pace}/20min | tz: {co.get('timezone')}")
    for k in ("no_phone", "dup_in_list", "opted_out", "already_enrolled",
              "generic_greeting", "to_enroll"):
        print(f"  {k}: {counts[k]}")
    if schedule:
        print(f"  first send due: {schedule[0]} | last: {schedule[-1]}")
    for p in to_enroll[:5]:
        print(f"    sample: {p['name']!r} first_name={p['first_name']!r} "
              f"{p['phone']} {p['email'] or ''}")

    if args.dry_run:
        print("DRY RUN — nothing written.")
        return
    if not to_enroll:
        print("Nothing to enroll.")
        return

    # ---- write contacts (new ones only; existing rows left untouched) ----
    # first_name keeps the literal 'there' for business-named contacts: the
    # dispatcher's fallback is name's first token ("Hey 136-146,"), so the
    # graceful greeting must be stored explicitly.
    new_rows = [{
        "client_id": company_id, "name": p["name"],
        "first_name": p["first_name"], "last_name": p["last_name"],
        "phone": p["phone"], "email": p["email"],
        "city": p["city"], "state": p["state"], "type": "customer",
        "tags": ["Review Requested"],
    } for p in to_enroll if not p["existing_contact_id"]]

    created = sb.insert("contacts", new_rows, upsert_on="client_id,phone",
                        returning=True) if new_rows else []
    id_by_phone = {c["phone"]: c["id"] for c in created}

    requests_rows, slugs = [], set()
    for p, due in zip(to_enroll, schedule):
        cid = p["existing_contact_id"] or id_by_phone.get(p["phone"])
        if not cid:
            die(f"no contact id came back for {p['phone']} — aborting before "
                f"review_requests insert")
        slug = uuid.uuid4().hex[:12]
        while slug in slugs:
            slug = uuid.uuid4().hex[:12]
        slugs.add(slug)
        requests_rows.append({
            "company_id": company_id, "contact_id": cid,
            "campaign_type": "reactivation", "tracking_slug": slug,
            "status": "staged" if args.staged else "pending",
            "step_number": 1, "next_send_at": due,
        })
    sb.insert("review_requests", requests_rows)
    verb = "STAGED (not active)" if args.staged else "ENROLLED"
    print(f"{verb} {len(requests_rows)} review_requests "
          f"({len(created)} new contacts, "
          f"{len(to_enroll) - len(new_rows)} existing reused).")


if __name__ == "__main__":
    main()
